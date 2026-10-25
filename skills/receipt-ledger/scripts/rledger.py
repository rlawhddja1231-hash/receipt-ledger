#!/usr/bin/env python3
"""rledger.py — 영수증 장부 도구 (Claude가 영수증·세금계산서·거래명세서·카드 문자를 엑셀 장부로 만드는 도구)

순서
  1) python3 rledger.py prep  영수증폴더/ 파일.pdf 사진.jpg [--out 작업]    # 읽을 준비: PDF 글자 꺼내기, 쪽·사진을 읽기 좋은 그림으로
       → 작업/작업목록.md (파일마다 할 일) + 작업/입력.json (글에서 찾은 값으로 채운 초안)
  2) Claude가 작업/그림의 사진을 직접 보고 입력.json 의 빈칸·틀린 값을 고친다
  3) python3 rledger.py check 작업/입력.json                                  # 검사: 사업자번호, 공급가+부가세=합계, 날짜, 중복…
  4) python3 rledger.py build 작업/입력.json --out 장부.xlsx [--into 지난장부.xlsx]  # 엑셀 장부(장부·월별·거래처별·계정과목별·확인 필요)

그 밖에
  python3 rledger.py sms   카드문자.txt [--year 2026] [--out 입력.json]   # 카드 승인 문자 → 입력.json
  python3 rledger.py parse 글.txt                                            # 글에서 날짜·상호·사업자번호·금액 찾기(시험용)

준비물: pip install --break-system-packages openpyxl pypdf pypdfium2 pillow
"""
import argparse
import datetime as dt
import json
import os
import re
import sys

TODAY = dt.date.today()
FIELDS = ["파일", "종류", "날짜", "상호", "사업자번호", "품목", "공급가액", "부가세", "합계", "결제수단", "카드끝4자리",
          "승인번호", "계정과목", "메모", "확신"]
KINDS = ["영수증", "카드영수증", "현금영수증", "세금계산서", "계산서", "거래명세서", "카드문자"]
IMG_EXT = (".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif", ".tif", ".tiff", ".heic", ".heif")


def _need(mod, pkg=None):
    try:
        return __import__(mod)
    except ImportError:
        print(f"{pkg or mod} 가 없습니다: pip install --break-system-packages {pkg or mod}", file=sys.stderr)
        sys.exit(2)


# ───────────────────────── 값 다루기 ─────────────────────────

def digits(s):
    return re.sub(r"\D", "", str(s or ""))


def biz_ok(s):
    d = digits(s)
    if len(d) != 10:
        return False
    w = [1, 3, 7, 1, 3, 7, 1, 3, 5]
    n = [int(x) for x in d]
    t = sum(a * b for a, b in zip(n[:9], w)) + (n[8] * 5) // 10
    return (10 - t % 10) % 10 == n[9]


def biz_fmt(s):
    d = digits(s)
    return f"{d[:3]}-{d[3:5]}-{d[5:]}" if len(d) == 10 else (str(s).strip() if s else "")


MINUS = "-−–—△▲"


def money(v):
    """'12,000원', '₩12,000', 12000, '-3,000', '₩-3,000', '−22,000', '(1,000)', '1.100.000' → int. 못 읽으면 None."""
    if v is None or v == "" or isinstance(v, (bool, dict, list, tuple)):
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v != v:  # NaN
            return None
        return int(round(v))
    s = str(v).strip().replace(" ", "")
    s = re.sub(r"^(일금|금)", "", s)
    s = re.sub(r"(원정|원|₩|\\|KRW|정)", "", s)
    neg = any(s.startswith(c) for c in MINUS) or (s.startswith("(") and s.endswith(")"))
    s = s.strip("()" + MINUS)
    if re.fullmatch(r"\d{1,3}(\.\d{3}){2,}", s):  # 1.100.000 (점을 쉼표처럼 쓴 것)
        s = s.replace(".", "")
    s = s.replace(",", "")
    if not re.fullmatch(r"\d+(\.\d+)?", s):
        return None
    n = int(round(float(s)))
    return -n if neg else n


def to_date(v, year=None):
    if v in (None, "") or isinstance(v, (bool, dict, list)):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    m = re.search(r"(\d{4})\s*[-./년]\s*(\d{1,2})\s*[-./월]\s*(\d{1,2})", s)
    guess_year = False
    if not m:
        m2 = re.search(r"(?<!\d)(\d{2})[-./](\d{1,2})[-./](\d{1,2})(?!\d)", s)
        if m2:
            y = 2000 + int(m2.group(1))
            mo, d = int(m2.group(2)), int(m2.group(3))
        else:
            m3 = re.search(r"(?<!\d)(\d{4})(\d{2})(\d{2})(?!\d)", s)
            if m3:
                y, mo, d = int(m3.group(1)), int(m3.group(2)), int(m3.group(3))
            else:
                m4 = re.search(r"(?<!\d)(\d{1,2})\s*[/.월]\s*(\d{1,2})일?(?!\d)", s)
                if not m4:
                    return None
                y, mo, d = (year or TODAY.year), int(m4.group(1)), int(m4.group(2))
                guess_year = year is None
    else:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
    try:
        out = dt.date(y, mo, d)
    except ValueError:
        return None
    if guess_year and out > TODAY + dt.timedelta(days=7):  # 연도 없는 12/30 을 10월에 읽으면 작년
        try:
            out = out.replace(year=out.year - 1)
        except ValueError:
            return None
    return out


CARD_FULL = re.compile(r"(?<!\d)(\d{4})[- ]?(\d{4}|\d{2}\*{2}|\*{4})[- ]?(\d{4}|\*{4})[- ]?(\d{3,4})(?!\d)|(?<!\d)\d{15,16}(?!\d)")


def mask_card(s):
    """카드번호는 끝 4자리만 남긴다."""
    d = digits(s)
    return d[-4:] if len(d) >= 4 else d


def mask_cards_in(text):
    """글 속 카드번호(16자리)를 ****-****-****-1234 로."""
    if not isinstance(text, str):
        return text
    return CARD_FULL.sub(lambda m: "****-****-****-" + digits(m.group(0))[-4:], text)


# ───────────────────────── 글에서 값 찾기 ─────────────────────────

LBL_SUPPLY = r"(?:공급\s*가\s*액|공급\s*가|과세\s*물품\s*가액|과세\s*금액|과세\s*물품|물품\s*가액|공급\s*금액)"
LBL_VAT = r"(?:부\s*가\s*세|부가\s*가치\s*세|세\s*액|VAT|V\.A\.T)"
LBL_TOTAL = r"(?:합\s*계\s*금\s*액|합\s*계|총\s*액|총\s*합\s*계|결\s*제\s*금\s*액|받\s*을\s*금\s*액|청\s*구\s*금\s*액|총\s*결제\s*금액|판매\s*총액|TOTAL|Total|총\s*금액|승인\s*금액)"
LBL_NAME = r"(?:상\s*호\s*명?|가\s*맹\s*점\s*명?|매\s*장\s*명|점\s*포\s*명|업\s*체\s*명|법\s*인\s*명)"
LBL_DATE = r"(?:작성\s*일\s*자|거래\s*일\s*시|거래\s*일\s*자|승인\s*일\s*시|판매\s*일\s*시|발행\s*일\s*자|일\s*시|날\s*짜|거래일)"
NUM = r"([-−]?\s*[₩\\]?\s*[-−]?\d{1,3}(?:,\d{3})+|[-−]?\s*[₩\\]?\s*[-−]?\d+)\s*원?"
NOT_TOTAL = re.compile(r"포인트|누적|적립|잔액|할인|마일리지|받은\s*금액|거스름|면세\s*물품|봉사료")


def _amount_after(label, text, skip=None):
    vals = []
    for line in text.splitlines():
        if skip is not None and skip.search(line):
            continue
        m = re.search(label + r"\s*[:：]?\s*(?:\(원\))?\s*" + NUM, line)
        if m:
            v = money(m.group(1))
            if v is not None:
                vals.append(v)
    return vals


def _table_amounts(text):
    """'작성일자 공급가액 세액' 다음 줄 '2026-10-20 1,500,000 150,000' 같은 표 모양."""
    lines = [x for x in text.splitlines() if x.strip()]
    sup, vat, tot = [], [], []
    for i, line in enumerate(lines[:-1]):
        heads = re.findall(LBL_SUPPLY + "|" + LBL_VAT + "|" + LBL_TOTAL + r"|작성\s*일자|품\s*목|수\s*량|단\s*가|월\s*일", line)
        if len(heads) < 2:
            continue
        nxt = lines[i + 1]
        nums = re.findall(r"[-−]?\d{1,3}(?:,\d{3})+|\d{4,}-\d{2}-\d{2}|[-−]?\d+", nxt)
        nums = [n for n in nums if not re.fullmatch(r"\d{4,}-\d{2}-\d{2}", n)]
        heads_amt = [h for h in heads if re.fullmatch(LBL_SUPPLY + "|" + LBL_VAT + "|" + LBL_TOTAL, h)]
        if len(nums) < len(heads_amt) or not heads_amt:
            continue
        vals = nums[-len(heads_amt):]
        for h, v in zip(heads_amt, vals):
            (sup if re.fullmatch(LBL_SUPPLY, h) else vat if re.fullmatch(LBL_VAT, h) else tot).append(money(v))
        break
    return sup, vat, tot


ROLE_RE = re.compile(r"공\s*급\s*받\s*는\s*자|공\s*급\s*자|매\s*입\s*자|매\s*출\s*자")
BIZ_RE = re.compile(r"(?<!\d)(\d{3})\s*-\s*(\d{2})\s*-\s*(\d{5})(?!\d)")


def _seller_fields(t):
    """세금계산서·계산서에서 '공급자'(파는 쪽)의 사업자번호와 상호를 고른다.
    글 순서대로 '공급자/공급받는자' 이름표를 따라가며, 한 줄에 값이 둘이면 이름표가 처음 나온 순서로 나눈다."""
    labels = [("buyer" if re.search("받", m.group(0)) or "매입" in m.group(0) else "seller", m.start())
              for m in ROLE_RE.finditer(t)]
    if not labels:
        return None, None
    first_role = labels[0][0]

    def role_at(pos):
        r = None
        for role, p in labels:
            if p <= pos:
                r = role
        return r

    biz, name = None, None
    for line_m in re.finditer(r"[^\n]+", t):
        line, base = line_m.group(0), line_m.start()
        bs = list(BIZ_RE.finditer(line))
        names = list(re.finditer(LBL_NAME + r"\s*[:：]?\s*(?:\(\s*법인명\s*\)\s*)?([^\s:：][^\n]*?)(?=\s+" + LBL_NAME + r"|\s{2,}|\s*(?:성\s*명|대표|$))", line))
        for group, kind in ((bs, "biz"), (names, "name")):
            if not group:
                continue
            if len(group) >= 2:  # 한 줄에 둘: 이름표가 처음 나온 순서대로 (보통 왼쪽이 공급자)
                pick = group[0] if first_role == "seller" else group[1]
            else:
                pick = group[0] if role_at(base + group[0].start()) in ("seller", None) else None
            if pick is None:
                continue
            if kind == "biz" and biz is None:
                biz = "-".join(pick.groups())
            elif kind == "name" and name is None:
                nm = pick.group(1).strip(" :：")
                nm = re.sub(r"^\(\s*법인명\s*\)\s*", "", nm)
                if nm and not re.fullmatch(r"[\d\-\s]+", nm):
                    name = nm[:40]
    return biz, name


def parse_text(text, year=None):
    """영수증·세금계산서 글 → 초안 dict (못 찾은 칸은 빈칸)."""
    t = text.replace("　", " ")
    out = {k: "" for k in FIELDS}
    if re.search(r"세\s*금\s*계\s*산\s*서", t):
        out["종류"] = "세금계산서"
    elif re.search(r"(?<!세금)계\s*산\s*서", t) and re.search(r"면\s*세", t):
        out["종류"] = "계산서"
    elif "현금영수증" in t.replace(" ", "") or "현금(소득공제)" in t or "현금(지출증빙)" in t:
        out["종류"] = "현금영수증"
    elif re.search(r"거\s*래\s*명\s*세", t):
        out["종류"] = "거래명세서"
    elif re.search(r"카드|승인\s*번호|CARD", t, re.I):
        out["종류"] = "카드영수증"
    elif t.strip():
        out["종류"] = "영수증"
    d = None
    for line in t.splitlines():
        if re.search(LBL_DATE, line):
            d = to_date(line, year)
            if d:
                break
    if not d:
        d = to_date(t, year)
    out["날짜"] = d.isoformat() if d else ""
    sb, sn = (None, None)
    if out["종류"] in ("세금계산서", "계산서", "거래명세서"):
        sb, sn = _seller_fields(t)
    bizs = BIZ_RE.findall(t)
    out["사업자번호"] = sb or ("-".join(bizs[0]) if bizs else "")
    if sn:
        out["상호"] = sn
    else:
        for line in t.splitlines():
            m = re.search(LBL_NAME + r"\s*[:：]?\s*(.+)", line)
            if m:
                name = re.split(r"\s{2,}|\s*(?:대표자|대표|성명|사업자|등록번호|상\s*호|TEL|전화|☎)", m.group(1).strip())[0].strip(" :：")
                name = re.sub(r"^\(\s*법인명\s*\)\s*", "", name)
                if name and not re.fullmatch(r"[\d\-\s]+", name):
                    out["상호"] = name[:40]
                    break
    sup = _amount_after(LBL_SUPPLY, t)
    vat = _amount_after(LBL_VAT, t)
    tot = _amount_after(LBL_TOTAL, t, skip=NOT_TOTAL)
    ts, tv, tt = _table_amounts(t)
    sup, vat, tot = sup or ts, vat or tv, tot or tt
    tot = [x for x in tot if x is not None]
    if sup:
        out["공급가액"] = sup[-1]
    if vat:
        out["부가세"] = vat[-1]
    if tot:
        # 공급가액+부가세와 맞는 합계가 있으면 그것, 없으면 마지막에 나온 합계
        s, v = out["공급가액"], out["부가세"]
        match = [x for x in tot if s != "" and v != "" and abs(s + v - x) <= 1]
        out["합계"] = match[-1] if match else tot[-1]
    if out["종류"] == "계산서" and out["부가세"] == "":
        out["부가세"] = 0
    m = re.search(r"승\s*인\s*번\s*호\s*[:：]?\s*(\d{8}-\d{8}-\d{8}|\d{6,12})", t)
    if m:
        out["승인번호"] = m.group(1)
    m = CARD_FULL.search(t)
    if m and not m.group(0).replace("-", "").replace(" ", "")[-4:].startswith("*"):
        out["카드끝4자리"] = mask_card(m.group(0))
    if out["종류"] in ("카드영수증",) or re.search(r"신용|체크|카드", t):
        out["결제수단"] = "카드"
    elif re.search(r"현\s*금", t):
        out["결제수단"] = "현금"
    elif re.search(r"계\s*좌|이\s*체|송\s*금", t):
        out["결제수단"] = "계좌이체"
    return out


CARD_CO = r"(?:신한|KB국민|KB|국민|삼성|현대|롯데|하나|우리|BC|비씨|NH농협|NH|농협|씨티|IBK|기업|카카오뱅크|카카오페이|토스|케이뱅크|광주|전북|제주|수협|신협|우체국)"
SMS_START = re.compile(r"\[Web발신\]|" + CARD_CO + r"\s*(?:카드|체크)?\s*\(?\d{4}\)?\s*(?:승인|취소|해외|일시불|할부)|"
                       + CARD_CO + r"\s*(?:카드|체크)\s*(?:승인|취소|해외)")


def _sms_blocks(text):
    """문자 모음 → 문자 하나씩. 빈 줄·[Web발신]·카드사 이름으로 시작하는 새 문자에서 나눈다."""
    blocks, cur = [], []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            if cur:
                blocks.append(" ".join(cur))
                cur = []
            continue
        starts = SMS_START.search(line)
        if starts and starts.start() == 0 and cur and re.search(r"\d\s*원|USD|\$", " ".join(cur)):
            blocks.append(" ".join(cur))
            cur = []
        cur.append(line)
    if cur:
        blocks.append(" ".join(cur))
    return blocks


def parse_sms(text, year=None, report=None):
    """카드 승인 문자들 → 초안 목록. 한 문자에 한 건. report 리스트에 못 읽은 문자를 담는다."""
    out = []
    for b in _sms_blocks(text):
        one = re.sub(r"\s+", " ", b.replace("[Web발신]", " ")).strip()
        if not one:
            continue
        mc = re.search(CARD_CO, one)
        if not mc or not re.search(r"승인|취소|일시불|할부|사용|결제", one):
            if report is not None and re.search(r"\d\s*원|USD|\$", one):
                report.append(one[:60])
            continue
        last4 = re.search(r"(?:카드|체크)?\s*\(?(\d{4})\)?\s*(?:승인|취소|해외|일시불|할부|사용)", one[mc.start():mc.start() + 30])
        amt = re.search(r"(?<![\d.])(\d{1,3}(?:,\d{3})+|\d+)\s*원", one)
        fx = re.search(r"(USD|US\$|JPY|EUR|CNY|\$)\s*([\d,]+(?:\.\d+)?)|([\d,]+(?:\.\d+)?)\s*(USD|달러|엔|유로)", one)
        md = re.search(r"(?<!\d)(\d{1,2})/(\d{1,2})(?!\d)", one)
        hm = re.search(r"(?<!\d)(\d{1,2}):(\d{2})(?!\d)", one)
        if not (amt or fx) or not md:
            if report is not None:
                report.append(one[:60])
            continue
        e = {k: "" for k in FIELDS}
        d = to_date(f"{md.group(1)}/{md.group(2)}", year)
        cancel = "취소" in one
        memo = [mc.group(0) + "카드"]
        if cancel:
            memo.append("취소")
        inst = re.search(r"(\d{1,2})\s*개월", one)
        if inst:
            memo.append(f"할부 {inst.group(1)}개월")
        total = money(amt.group(1)) if amt else ""
        if fx and not amt:
            cur = fx.group(1) or fx.group(4)
            val = fx.group(2) or fx.group(3)
            memo.append(f"해외승인 {cur} {val} — 원화 금액은 카드 청구서로 확인")
        # 상호: 금액·날짜·시간·할부 글자 뒤에 남은 글
        ends = [x.end() for x in (amt, fx, md, hm) if x]
        tail = one[max(ends):] if ends else ""
        tail = re.sub(r"^[\s/]*(?:\(?\s*(?:일시불|\d{1,2}\s*개월|할부)\s*\)?[\s/]*)+", "", tail)
        tail = re.sub(r"\s*(?:누적|잔액|총\s*누적|총|사용\s*가능|한도).*$", "", tail).strip(" ,.-/")
        if not tail:  # 상호가 금액 앞에 있는 모양
            head = one[mc.end():min(x.start() for x in (amt, fx, md) if x)]
            head = re.sub(r"\(?\d{4}\)?|승인|취소|해외|일시불|할부|카드|체크|[가-힣]\*[가-힣]?|님", " ", head)
            tail = re.sub(r"\s+", " ", head).strip(" ,.-")
        e.update({"종류": "카드문자", "날짜": d.isoformat() if d else "", "결제수단": "카드",
                  "카드끝4자리": last4.group(1) if last4 else "", "합계": (-abs(total) if cancel and total != "" else total),
                  "상호": tail[:40], "메모": ", ".join(memo), "확신": "보통"})
        out.append(e)
    return out


# ───────────────────────── 계정과목 짐작 ─────────────────────────

ACCOUNTS = [
    ("여비교통비", r"택시|카카오T|티머니|코레일|KTX|SRT|고속버스|버스|지하철|주차|하이패스|도로공사|톨게이트|항공|대한항공|아시아나|제주항공"),
    ("차량유지비", r"주유|오일|SK에너지|GS칼텍스|S-?OIL|현대오일뱅크|알뜰주유|세차|정비|타이어|카센터|충전소"),
    ("통신비", r"KT|SKT|SK텔레콤|LG\s?U\+|유플러스|인터넷|통신|알뜰폰"),
    ("소모품비", r"다이소|문구|오피스|알파|모닝글로리|소모품|토너|복사지|쿠팡|철물"),
    ("도서인쇄비", r"서점|교보|영풍|알라딘|예스24|YES24|인쇄|출력|제본|킨코스"),
    ("지급수수료", r"수수료|등기|우체국|택배|CJ대한통운|한진|로젠|퀵"),
    ("수도광열비", r"전기|한국전력|한전|가스|도시가스|수도"),
    ("복리후생비(식대)", r"식당|김밥|분식|국밥|치킨|피자|버거|맥도날드|롯데리아|배달의민족|요기요|쿠팡이츠|한식|중식|일식|"
                      r"카페|커피|스타벅스|이디야|투썸|메가커피|빽다방|베이커리|파리바게뜨|뚜레쥬르|편의점|GS25|CU|세븐일레븐|이마트24|식사|도시락|음식"),
    ("비품", r"하이마트|전자랜드|삼성전자|LG전자|애플|컴퓨터|모니터|노트북|의자|책상|가구|이케아"),
    ("광고선전비", r"광고|네이버\s?광고|카카오\s?광고|현수막|전단|명함"),
    ("교육훈련비", r"교육|학원|강의|세미나|클래스101|인프런"),
]


def guess_account(e):
    hay = f"{e.get('상호', '')} {e.get('품목', '')}"
    for name, rx in ACCOUNTS:
        if re.search(rx, hay, re.I):
            return name
    return ""


# ───────────────────────── 검사 ─────────────────────────

def normalize(e, year=None):
    """값 모양을 맞추고 비어 있는 금액을 채운다. 바꾼 것은 '계산' 목록에 남긴다."""
    e = {k: e.get(k, "") for k in FIELDS} | {k: v for k, v in e.items() if k not in FIELDS}
    calc = []
    for k in ("파일", "종류", "상호", "사업자번호", "품목", "결제수단", "카드끝4자리", "승인번호", "계정과목", "메모", "확신"):
        v = e[k]
        if v is None or (isinstance(v, float) and v != v):
            v = ""
        elif isinstance(v, (list, tuple)):
            v = ", ".join(str(x) for x in v if x not in (None, ""))
        elif isinstance(v, dict):
            v = ", ".join(f"{a}: {b}" for a, b in v.items())
        elif isinstance(v, float) and v.is_integer():
            v = str(int(v))
        e[k] = mask_cards_in(str(v).strip()) if v != "" else ""
    d = to_date(e["날짜"], year)
    e["날짜"] = d.isoformat() if d else (str(e["날짜"]) if e["날짜"] else "")
    if e["사업자번호"]:
        e["사업자번호"] = biz_fmt(e["사업자번호"])
    for k in ("공급가액", "부가세", "합계"):
        v = money(e[k])
        e[k] = "" if v is None else v
    if e["카드끝4자리"]:
        e["카드끝4자리"] = mask_card(e["카드끝4자리"])
    tax_free = e["종류"] == "계산서" or "면세" in str(e.get("메모", ""))
    s, v, t = e["공급가액"], e["부가세"], e["합계"]
    if t == "" and s != "" and v != "":
        e["합계"] = s + v
        calc.append("합계=공급가액+부가세로 계산")
    elif s == "" and t != "" and v != "":
        e["공급가액"] = t - v
        calc.append("공급가액=합계-부가세로 계산")
    elif s == "" and v == "" and t != "":
        if tax_free:
            e["공급가액"], e["부가세"] = t, 0
            calc.append("면세로 보고 부가세 0")
        else:
            e["공급가액"] = round(t / 1.1)
            e["부가세"] = t - e["공급가액"]
            calc.append("합계만 있어 부가세 10%로 나눠 계산(추정)")
            if e["종류"] in ("세금계산서", "거래명세서"):
                e["_추정세금"] = True
    elif v == "" and s != "" and t != "":
        e["부가세"] = t - s
        calc.append("부가세=합계-공급가액으로 계산")
    if not e["계정과목"]:
        g = guess_account(e)
        if g:
            e["계정과목"] = g
            calc.append(f"계정과목 '{g}'은 상호로 짐작")
    e["_계산"] = calc
    return e


def problems(e, seen=None):
    """(틀림 목록, 확인 목록)"""
    bad, warn = [], []
    if not e["날짜"]:
        bad.append("날짜 없음")
    else:
        d = to_date(e["날짜"])
        if not d:
            bad.append(f"날짜로 못 읽음({e['날짜']})")
        elif d > TODAY:
            bad.append("날짜가 오늘보다 뒤")
        elif d < TODAY - dt.timedelta(days=365 * 6):
            warn.append("6년 넘은 날짜")
    if not e["상호"]:
        warn.append("상호 없음")
    if e["합계"] == "":
        bad.append("합계 없음")
    if e["사업자번호"]:
        if len(digits(e["사업자번호"])) != 10:
            bad.append("사업자번호 자리 수가 10자리가 아님")
        elif not biz_ok(e["사업자번호"]):
            bad.append("사업자번호 검사 숫자 틀림(잘못 읽었을 가능성)")
    elif e["종류"] in ("세금계산서", "계산서"):
        bad.append(f"{e['종류']}인데 사업자번호 없음")
    s, v, t = e["공급가액"], e["부가세"], e["합계"]
    if s != "" and v != "" and t != "" and abs(s + v - t) > 1:
        bad.append(f"공급가액+부가세({s + v:,})가 합계({t:,})와 다름")
    if s not in ("", 0) and v != "" and e["종류"] not in ("계산서",) and abs(t if t != "" else 0) > 0:
        expect = s * 0.1
        if v != 0 and abs(v - expect) > max(2, abs(expect) * 0.01):
            warn.append(f"부가세가 공급가액의 10%({round(expect):,})와 다름")
        if v == 0 and e["종류"] == "세금계산서":
            warn.append("세금계산서인데 부가세 0(영세율?)")
    if e.get("_추정세금"):
        warn.append(f"{e['종류']}에 공급가액·부가세가 적혀 있을 텐데 비어 있음 — 그림에서 확인(지금 값은 10%로 나눈 추정)")
    if e["종류"] and e["종류"] not in KINDS:
        warn.append(f"종류 '{e['종류']}'는 정해진 이름이 아님")
    if str(e.get("확신", "")).strip() in ("낮음", "low"):
        warn.append("사진이 흐려 읽은 값이 불확실")
    if seen is not None:
        keys = []
        if e["승인번호"]:
            keys.append(("승인", e["승인번호"], e["합계"]))
        if e["사업자번호"]:
            keys.append(("번호", e["날짜"], e["합계"], digits(e["사업자번호"])))
        if e["상호"]:
            keys.append(("상호", e["날짜"], e["합계"], re.sub(r"\s|\(주\)|주식회사", "", e["상호"])))
        if e["카드끝4자리"]:
            keys.append(("카드", e["날짜"], e["합계"], e["카드끝4자리"]))
        label = f"'{os.path.basename(str(e.get('파일') or ''))}'" if e.get("파일") else f"{e.get('_no', '?')}번"
        for k in keys:
            if k in seen:
                other_label, other_appr = seen[k]
                if e["승인번호"] and other_appr and e["승인번호"] != other_appr:
                    continue  # 승인번호가 다르면 같은 날 같은 금액이라도 다른 거래
                warn.append(f"{other_label}와 같은 거래로 보임(중복?)")
                break
        for k in keys:
            if k not in seen:
                seen[k] = (label, e["승인번호"])
    return bad, warn


def load_entries(path, year=None):
    if not os.path.exists(path):
        raise SystemExit(f"파일이 없습니다: {path}")
    raw = open(path, "rb").read()
    text = None
    for enc in ("utf-8-sig", "utf-16", "cp949"):
        try:
            text = raw.decode(enc)
            if enc == "utf-16" and raw[:2] not in (b"\xff\xfe", b"\xfe\xff"):
                continue
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise SystemExit("입력 파일의 글자 모양(인코딩)을 알 수 없습니다. UTF-8로 저장해 주세요.")
    try:
        data = json.loads(text, parse_constant=lambda c: None)
    except json.JSONDecodeError as err:
        raise SystemExit(f"입력 파일의 JSON 모양이 틀렸습니다({err.lineno}번째 줄): {err.msg}")
    if isinstance(data, dict):
        if "거래" in data or "entries" in data:
            data = data.get("거래", data.get("entries"))
        else:
            data = [data]
    if not isinstance(data, list):
        raise SystemExit("입력 파일은 거래 목록([ {...}, {...} ]) 모양이어야 합니다.")
    out = []
    for i, e in enumerate(data, 1):
        if not isinstance(e, dict):
            continue
        n = normalize(e, year)
        n["_no"] = i
        out.append(n)
    if not out:
        print("※ 입력 파일에 거래가 없습니다.")
    return out


def cmd_check(a):
    es = load_entries(a.file, a.year)
    seen = {}
    nb = nw = 0
    for e in es:
        bad, warn = problems(e, seen)
        if bad or warn or e["_계산"]:
            print(f"{e['_no']}번 {e['날짜'] or '?'} {e['상호'] or '?'} {e['합계'] if e['합계'] != '' else '?'}원")
            for x in bad:
                print(f"   ✗ {x}")
            for x in warn:
                print(f"   ! {x}")
            for x in e["_계산"]:
                print(f"   · {x}")
        nb += bool(bad)
        nw += bool(warn) and not bad
    ok = len(es) - nb - nw
    print(f"\n모두 {len(es)}건: 문제 없음 {ok}, 확인 필요 {nw}, 틀림 {nb}")
    tot = sum(e["합계"] for e in es if e["합계"] != "")
    print(f"합계 {tot:,}원 (공급가액 {sum(e['공급가액'] for e in es if e['공급가액'] != ''):,} / "
          f"부가세 {sum(e['부가세'] for e in es if e['부가세'] != ''):,})")


# ───────────────────────── 장부 만들기 ─────────────────────────

LEDGER_COLS = ["번호", "날짜", "종류", "상호", "사업자번호", "품목", "공급가액", "부가세", "합계", "결제수단", "카드끝4자리",
               "승인번호", "계정과목", "확인", "파일", "메모"]


def read_ledger(path):
    """지난번에 만든 장부.xlsx 의 '장부' 시트를 다시 읽는다(이어 쓰기용)."""
    openpyxl = _need("openpyxl")
    if not os.path.exists(path):
        raise SystemExit(f"지난 장부 파일이 없습니다: {path}")
    if not path.lower().endswith((".xlsx", ".xlsm")):
        raise SystemExit("지난 장부는 이 도구로 만든 엑셀(.xlsx)이어야 합니다.")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        wbf = openpyxl.load_workbook(path)
    except Exception as err:  # noqa: BLE001
        raise SystemExit(f"지난 장부를 열 수 없습니다: {err}")
    if "장부" not in wb.sheetnames:
        raise SystemExit(f"'{path}' 에 '장부' 시트가 없습니다. 이 도구로 만든 장부만 이어 쓸 수 있습니다.")
    ws, wsf = wb["장부"], wbf["장부"]
    head = [str(c.value).strip() if c.value is not None else None for c in ws[1]]
    if "날짜" not in head or "합계" not in head:
        raise SystemExit(f"'{path}' 의 '장부' 시트 첫 줄에 '날짜'·'합계' 칸이 없습니다. 칸 이름을 바꾸지 않은 장부만 이어 쓸 수 있습니다.")
    total_words = re.compile(r"^\s*(합\s*계|소\s*계|총\s*계|누\s*계|총합|total|sum)\s*$", re.I)
    out = []
    for ri, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        r = dict(zip(head, row))
        if any(isinstance(v, str) and total_words.match(v) for v in row):
            continue  # 합계·총계 줄
        if any(isinstance(c.value, str) and re.match(r"=\s*SUM", c.value, re.I) for c in wsf[ri]):
            continue
        for k in ("공급가액", "부가세", "합계"):  # 사람이 넣은 계산식(=2000+2500)은 저장된 값, 없으면 직접 계산
            if k in head and r.get(k) is None:
                fv = wsf.cell(ri, head.index(k) + 1).value
                if isinstance(fv, str) and re.fullmatch(r"=[\d+\-*/(). ]+", fv):
                    try:
                        r[k] = eval(fv[1:], {"__builtins__": {}}, {})  # 숫자와 + - * / 만
                    except Exception:  # noqa: BLE001
                        pass
        if not r.get("날짜") and r.get("합계") in (None, ""):
            continue
        e = {k: ("" if r.get(k) is None else r.get(k)) for k in FIELDS if k in r}
        if isinstance(e.get("날짜"), (dt.date, dt.datetime)):
            e["날짜"] = (e["날짜"].date() if isinstance(e["날짜"], dt.datetime) else e["날짜"]).isoformat()
        e["_old"] = True
        out.append(e)
    return out


def cmd_build(a):
    openpyxl = _need("openpyxl")
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    new = load_entries(a.file, a.year)
    old = []
    if a.into:
        old = [normalize(e, a.year) for e in read_ledger(a.into)]
        if not old:
            print("※ 지난 장부에서 거래를 하나도 못 읽었습니다. 새 거래만으로 만듭니다.")
    allx = old + new
    allx.sort(key=lambda e: (e["날짜"] or "9999", e.get("_old") is not True))
    seen = {}
    rows, issues = [], []
    skipped = 0
    dropped = []
    for i, e in enumerate(allx, 1):
        e["_no"] = i
        bad, warn = problems(e, seen)
        dup = any("중복" in w for w in warn)
        e["_warn"] = warn
        if dup and not e.get("_old") and (a.into or a.drop_dupes):
            skipped += 1  # 이미 있는 거래(지난 장부·같은 영수증 두 장·카드 문자와 영수증)는 한 번만
            dropped.append(e)
            continue
        e["_bad"], e["_warn"] = bad, warn
        rows.append(e)
    for n, e in enumerate(rows, 1):
        e["_no"] = n
    for e in dropped:
        issues.append(("-", "참고", "중복으로 보여 장부에서 뺌 — " + "; ".join(w for w in e.get("_warn", []) if "중복" in w)
                       if e.get("_warn") else "중복으로 보여 장부에서 뺌", e))
    for e in rows:
        n = e["_no"]
        for x in e["_bad"]:
            issues.append((n, "틀림", x, e))
        for x in e["_warn"]:
            issues.append((n, "확인", x, e))
        for x in e["_계산"]:
            if "추정" in x or "짐작" in x:
                issues.append((n, "참고", x, e))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "장부"
    hfill = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")
    red = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
    yel = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
    ws.append(LEDGER_COLS)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), hfill
        c.alignment = Alignment(horizontal="center")
    for e in rows:
        mark = "틀림: " + "; ".join(e["_bad"]) if e["_bad"] else ("확인: " + "; ".join(e["_warn"]) if e["_warn"] else "")
        d = to_date(e["날짜"])
        ws.append([e["_no"], d or e["날짜"], e["종류"], e["상호"], e["사업자번호"], e["품목"],
                   e["공급가액"] if e["공급가액"] != "" else None, e["부가세"] if e["부가세"] != "" else None,
                   e["합계"] if e["합계"] != "" else None, e["결제수단"], str(e["카드끝4자리"] or ""),
                   str(e["승인번호"] or ""), e["계정과목"], mark, os.path.basename(str(e["파일"] or "")), e["메모"]])
        r = ws.max_row
        ws.cell(r, 2).number_format = "yyyy-mm-dd"
        for c in (7, 8, 9):
            ws.cell(r, c).number_format = "#,##0"
        if e["_bad"] or e["_warn"]:
            for c in range(1, len(LEDGER_COLS) + 1):
                ws.cell(r, c).fill = red if e["_bad"] else yel
    last = ws.max_row
    ws.append(["합계", None, None, f"{len(rows)}건", None, None, f"=SUM(G2:G{last})", f"=SUM(H2:H{last})", f"=SUM(I2:I{last})"])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
        if c.column in (7, 8, 9):
            c.number_format = "#,##0"
    widths = [6, 11, 11, 22, 13, 22, 12, 11, 12, 9, 9, 11, 16, 40, 20, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(LEDGER_COLS))}{last}"
    dv = DataValidation(type="list", formula1='"' + ",".join(KINDS) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C2:C{last + 200}")

    def summary(title, keyf):
        s = wb.create_sheet(title)
        groups = {}
        for e in rows:
            k = keyf(e)
            g = groups.setdefault(k, [0, 0, 0, 0])
            g[0] += 1
            g[1] += e["공급가액"] or 0
            g[2] += e["부가세"] or 0
            g[3] += e["합계"] or 0
        s.append([title.replace("별", ""), "건수", "공급가액", "부가세", "합계"])
        for c in s[1]:
            c.font, c.fill = Font(bold=True), hfill
        for k, g in sorted(groups.items(), key=lambda kv: (kv[0] == "(없음)", kv[0]) if title == "월별" else (-kv[1][3],)):
            s.append([k] + g)
        s.append(["합계", sum(g[0] for g in groups.values()), sum(g[1] for g in groups.values()),
                  sum(g[2] for g in groups.values()), sum(g[3] for g in groups.values())])
        for c in s[s.max_row]:
            c.font = Font(bold=True)
        for row in s.iter_rows(min_row=2, min_col=3, max_col=5):
            for c in row:
                c.number_format = "#,##0"
        s.column_dimensions["A"].width = 24
        for col in "BCDE":
            s.column_dimensions[col].width = 14
        s.freeze_panes = "A2"
        return s

    summary("월별", lambda e: (to_date(e["날짜"]).isoformat()[:7] if to_date(e["날짜"]) else "(날짜 없음·오류)"))
    summary("거래처별", lambda e: e["상호"] or "(없음)")
    summary("계정과목별", lambda e: e["계정과목"] or "(정하지 않음)")
    sv = wb.create_sheet("부가세")
    sv.append(["구분", "건수", "공급가액", "부가세", "설명"])
    for c in sv[1]:
        c.font, c.fill = Font(bold=True), hfill
    groups = [("세금계산서", lambda e: e["종류"] == "세금계산서"), ("계산서(면세)", lambda e: e["종류"] == "계산서"),
              ("카드·현금영수증", lambda e: e["종류"] in ("카드영수증", "현금영수증", "카드문자")),
              ("그 밖의 영수증·종류 없음", lambda e: e["종류"] not in ("세금계산서", "계산서", "카드영수증", "현금영수증", "카드문자"))]
    for name, f in groups:
        g = [e for e in rows if f(e)]
        sv.append([name, len(g), sum(e["공급가액"] or 0 for e in g), sum(e["부가세"] or 0 for e in g),
                   "증빙별로 나눠 본 합계입니다. 실제 공제 가능 여부는 세무사·홈택스 기준으로 확인하세요."])
    sv.append(["합계", len(rows), sum(e["공급가액"] or 0 for e in rows), sum(e["부가세"] or 0 for e in rows), ""])
    for c in sv[sv.max_row]:
        c.font = Font(bold=True)
    for row in sv.iter_rows(min_row=2, min_col=3, max_col=4):
        for c in row:
            c.number_format = "#,##0"
    sv.column_dimensions["A"].width, sv.column_dimensions["E"].width = 18, 70
    for col in "BCD":
        sv.column_dimensions[col].width = 14
    si = wb.create_sheet("확인 필요")
    si.append(["번호", "구분", "내용", "날짜", "상호", "합계", "파일"])
    for c in si[1]:
        c.font, c.fill = Font(bold=True), hfill
    for n, kind, msg, e in issues:
        si.append([n, kind, msg, e["날짜"], e["상호"], e["합계"] if e["합계"] != "" else None, os.path.basename(str(e["파일"] or ""))])
        if kind in ("틀림", "확인"):
            for c in si[si.max_row]:
                c.fill = red if kind == "틀림" else yel
    for row in si.iter_rows(min_row=2, min_col=6, max_col=6):
        for c in row:
            c.number_format = "#,##0"
    for col, w in zip("ABCDEFG", (6, 7, 60, 11, 22, 12, 22)):
        si.column_dimensions[col].width = w
    for sh in wb.worksheets:  # 인쇄하면 가로 한 장 폭에 맞게
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth, sh.page_setup.fitToHeight = 1, 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    out = a.out or os.path.join(os.getcwd(), "장부.xlsx")
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
    wb.save(out)
    nb = sum(1 for e in rows if e["_bad"])
    nw = sum(1 for e in rows if e["_warn"] and not e["_bad"])
    print(f"장부 {len(rows)}건" + (f"(지난 장부 {len(old)}건 + 새 {len(new) - skipped}건)" if a.into else "")
          + (f", 중복으로 보여 뺀 것 {skipped}건('확인 필요' 시트에 적어 둠)" if skipped else "") + f" → {out}")
    print(f"합계 {sum(e['합계'] or 0 for e in rows):,}원 | 틀림 {nb}건(빨강), 확인 필요 {nw}건(노랑) — '확인 필요' 시트에 모아 둠")


# ───────────────────────── 읽을 준비 ─────────────────────────

def _render_pdf(path, outdir, base, dpi=150, max_pages=20):
    pdfium = _need("pypdfium2")
    pdf = pdfium.PdfDocument(path)
    pngs = []
    for i in range(min(len(pdf), max_pages)):
        img = pdf[i].render(scale=dpi / 72).to_pil()
        fn = os.path.join(outdir, f"{base}_p{i + 1}.png")
        img.save(fn)
        pngs.append(fn)
    return pngs, len(pdf)


def _pdf_text(path):
    """→ (쪽마다 글 목록, 문제). 열기 암호가 걸렸으면 문제='암호'."""
    pypdf = _need("pypdf")
    try:
        r = pypdf.PdfReader(path)
        if r.is_encrypted:
            try:
                if not r.decrypt(""):
                    return [], "암호"
            except Exception:  # noqa: BLE001
                return [], "암호"
        return [(p.extract_text() or "") for p in r.pages], ""
    except Exception as err:  # noqa: BLE001
        msg = str(err)
        return [], ("암호" if "password" in msg.lower() or "encrypt" in msg.lower() else f"글자 꺼내기 실패: {msg}")


def _prep_image(path, outdir, base, enhance=True):
    _need("PIL", "pillow")
    from PIL import Image, ImageOps, ImageFilter
    Image.MAX_IMAGE_PIXELS = 400_000_000  # 2억 화소 휴대폰 사진까지
    ext = os.path.splitext(path)[1].lower()
    try:
        im = Image.open(path)
    except Exception as err:  # noqa: BLE001
        im = None
        if ext in (".heic", ".heif"):
            try:
                import pillow_heif
                pillow_heif.register_heif_opener()
                im = Image.open(path)
            except ImportError:
                return None, "아이폰 HEIC 사진입니다. pip install --break-system-packages pillow-heif 를 하거나 JPG로 바꿔 주세요."
            except Exception as err2:  # noqa: BLE001
                return None, f"그림을 열 수 없음: {err2}"
        else:
            return None, f"그림을 열 수 없음: {err}"
    try:
        im.draft("RGB", (4000, 4000))
        im = ImageOps.exif_transpose(im).convert("RGB")
    except Exception as err:  # noqa: BLE001
        return None, f"그림을 열 수 없음(파일이 깨졌을 수 있음): {err}"
    w, h = im.size
    scale = min(1.0, 2000 / max(w, h))
    if scale < 1:
        im = im.resize((int(w * scale), int(h * scale)))
    fn = os.path.join(outdir, f"{base}.png")
    im.save(fn)
    outs = [fn]
    if enhance:  # 흐린 감열지 영수증용: 흑백 + 대비 + 선명하게
        g = ImageOps.grayscale(im).filter(ImageFilter.MedianFilter(3))
        hist = g.histogram()
        paper = max(range(128, 256), key=lambda v: hist[v])  # 종이 밝기(가장 많은 밝은 값)
        lo, hi = max(0, paper - 70), max(paper - 3, 1)
        g = g.point([0 if v <= lo else 255 if v >= hi else int((v - lo) * 255 / max(1, hi - lo)) for v in range(256)])
        g = g.filter(ImageFilter.SHARPEN)
        fn2 = os.path.join(outdir, f"{base}_선명.png")
        g.save(fn2)
        outs.append(fn2)
    return outs, ""


def cmd_prep(a):
    files = []
    for p in a.paths:
        if os.path.isdir(p):
            for root, _, names in os.walk(p):
                for n in sorted(names):
                    if n.lower().endswith(IMG_EXT + (".pdf", ".txt")):
                        files.append(os.path.join(root, n))
        elif os.path.exists(p):
            files.append(p)
        else:
            print(f"없음: {p}")
    if not files:
        raise SystemExit("읽을 파일(사진·PDF)이 없습니다.")
    outdir = a.out or os.path.join(os.getcwd(), "영수증_작업")
    gdir = os.path.join(outdir, "그림")
    os.makedirs(gdir, exist_ok=True)
    entries, lines = [], ["# 영수증 읽기 작업 목록", "",
                          "각 파일의 그림을 직접 보고 입력.json 의 같은 번호 칸을 채우거나 고친다. 글에서 찾은 값(초안)도 그림과 견줘 본다.", ""]
    for i, f in enumerate(files, 1):
        base = f"{i:03d}_" + re.sub(r"[^\w가-힣.-]", "_", os.path.splitext(os.path.basename(f))[0])[:40]
        ext = os.path.splitext(f)[1].lower()
        draft = {k: "" for k in FIELDS}
        draft["파일"] = os.path.basename(f)
        note, pngs, text = "", [], ""
        if ext == ".pdf":
            pages, err = _pdf_text(f)
            text = "\n".join(pages)
            if err == "암호":
                note = "암호 걸린 PDF — 암호를 푼 파일이 필요"
            else:
                if err:
                    note = err
                try:
                    pngs, npages = _render_pdf(f, gdir, base, max_pages=a.max_pages)
                    if npages > a.max_pages:
                        note = f"{npages}쪽 중 앞 {a.max_pages}쪽만 그림으로(--max-pages 로 늘릴 수 있음)"
                except Exception as err2:  # noqa: BLE001
                    note = note or f"쪽 그림 실패: {err2}"
            # 쪽마다 따로 된 증빙(세금계산서 여러 장 묶음)이면 쪽마다 한 건
            per = [pg for pg in pages if len(re.sub(r"\s", "", pg)) >= 30 and
                   re.search(LBL_TOTAL + "|" + LBL_SUPPLY, pg) and to_date(pg, a.year)]
            if len(pages) >= 2 and len(per) == len(pages):
                lines.append(f"## {i}. {os.path.basename(f)} — 글이 있는 PDF, 쪽마다 따로 된 증빙 {len(pages)}건(쪽마다 그림과 견줘 확인)")
                for pi, pg in enumerate(pages, 1):
                    dd = {k: "" for k in FIELDS}
                    dd.update({k: v for k, v in parse_text(pg, a.year).items() if v not in ("", None)})
                    dd["파일"] = f"{os.path.basename(f)}#{pi}쪽"
                    dd["확신"] = "보통"
                    entries.append(dd)
                    lines.append(f"- {pi}쪽 그림: {pngs[pi - 1] if pi <= len(pngs) else '(그림 없음 — 쪽 수 제한)'}")
                tf = os.path.join(outdir, f"{base}_글.txt")
                open(tf, "w", encoding="utf-8").write(mask_cards_in("\n\n=== 쪽 나눔 ===\n\n".join(pages)))
                if note:
                    lines.append(f"- 주의: {note}")
                lines.append("")
                continue
        elif ext == ".txt":
            raw = open(f, "rb").read()
            for enc in ("utf-8-sig", "utf-16", "cp949"):
                try:
                    text = raw.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            missed = []
            sms = parse_sms(text, a.year, missed)
            if sms:  # 카드 승인 문자 모음이면 문자마다 한 건
                for e in sms:
                    e["파일"] = os.path.basename(f)
                entries.extend(sms)
                lines.append(f"## {i}. {os.path.basename(f)} — 카드 승인 문자 {len(sms)}건(자동으로 채움, 품목·사업자번호는 없음)")
                if missed:
                    lines.append(f"- 주의: 못 읽은 문자 {len(missed)}개 — 직접 보고 입력.json 에 더할 것")
                    lines += [f"  - {mask_cards_in(x)}" for x in missed]
                lines.append("")
                continue
        else:
            pngs, err = _prep_image(f, gdir, base, not a.no_enhance)
            pngs = pngs or []
            note = err
        has_text = len(re.sub(r"\s", "", text or "")) >= 30
        if has_text:
            draft.update({k: v for k, v in parse_text(text, a.year).items() if v not in ("", None) and k != "파일"})
            draft["확신"] = "보통"
            tf = os.path.join(outdir, f"{base}_글.txt")
            open(tf, "w", encoding="utf-8").write(mask_cards_in(text))  # 카드번호는 끝 4자리만
        entries.append(draft)
        if ext == ".pdf" and has_text:
            kind = "글이 있는 PDF(초안을 그림과 견줘 확인)"
        elif pngs:
            kind = "스캔·사진(그림을 보고 직접 읽기)"
        elif ext == ".txt":
            kind = "글 파일"
        else:
            kind = "읽을 수 없는 파일(아래 주의 참고) — 입력.json 에서 이 줄을 지우거나 다른 파일을 받을 것"
        lines.append(f"## {i}. {os.path.basename(f)} — {kind}")
        for pth in pngs:
            lines.append(f"- 그림: {pth}")
        if has_text:
            lines.append(f"- 꺼낸 글: {os.path.join(outdir, base + '_글.txt')}")
            filled = [k for k in ("날짜", "상호", "사업자번호", "공급가액", "부가세", "합계") if draft.get(k) not in ("", None)]
            lines.append(f"- 초안에서 찾은 칸: {', '.join(filled) or '없음'}")
        if note:
            lines.append(f"- 주의: {note}")
        lines.append("")
    jf = os.path.join(outdir, "입력.json")
    json.dump(entries, open(jf, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    md = os.path.join(outdir, "작업목록.md")
    open(md, "w", encoding="utf-8").write("\n".join(lines))
    print(f"파일 {len(files)}개 준비 → {outdir}")
    print(f"  작업목록: {md}\n  입력 초안: {jf}  (칸: {', '.join(FIELDS)})")
    n_text = sum(1 for e in entries if e.get("확신"))
    print(f"  거래 {len(entries)}건 — 글에서 초안을 채운 것 {n_text}건, 그림을 보고 읽어야 할 것 {len(entries) - n_text}건")
    warn = [ln for ln in lines if ln.startswith("- 주의")]
    if warn:
        print(f"  주의할 파일 {len(warn)}개 — 작업목록의 '주의'를 보세요")


def cmd_sms(a):
    text = open(a.file, encoding="utf-8", errors="replace").read()
    missed = []
    es = parse_sms(text, a.year, missed)
    if missed:
        print(f"※ 카드 문자처럼 보이지만 못 읽은 것 {len(missed)}개:")
        for x in missed:
            print(f"   {mask_cards_in(x)}")
    if not es:
        raise SystemExit("카드 승인 문자를 못 찾았습니다. 예: [Web발신] 신한카드(1234)승인 홍*동 12,000원 10/25 13:05 스타벅스")
    out = a.out or os.path.join(os.getcwd(), "입력_문자.json")
    json.dump(es, open(out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"카드 문자 {len(es)}건 → {out}")
    for e in es[:20]:
        amt = f"{e['합계']:,}원" if e["합계"] != "" else "원화 금액 없음"
        print(f"  {e['날짜']} {e['상호']} {amt} ({e['메모']}" + (f", 끝 {e['카드끝4자리']})" if e["카드끝4자리"] else ")"))


def cmd_parse(a):
    text = open(a.file, encoding="utf-8", errors="replace").read()
    print(json.dumps(parse_text(text, a.year), ensure_ascii=False, indent=1))


def main(argv=None):
    p = argparse.ArgumentParser(prog="rledger", description="영수증 장부 도구")
    sub = p.add_subparsers(dest="cmd", required=True)
    sp = sub.add_parser("prep"); sp.add_argument("paths", nargs="+"); sp.add_argument("--out"); sp.add_argument("--year", type=int)
    sp.add_argument("--no-enhance", action="store_true"); sp.add_argument("--max-pages", type=int, default=30)
    sp.set_defaults(fn=cmd_prep)
    sp = sub.add_parser("check"); sp.add_argument("file"); sp.add_argument("--year", type=int); sp.set_defaults(fn=cmd_check)
    sp = sub.add_parser("build"); sp.add_argument("file"); sp.add_argument("--out"); sp.add_argument("--into")
    sp.add_argument("--drop-dupes", action="store_true", help="중복으로 보이는 뒤쪽 거래는 장부에서 뺀다")
    sp.add_argument("--year", type=int); sp.set_defaults(fn=cmd_build)
    sp = sub.add_parser("sms"); sp.add_argument("file"); sp.add_argument("--out"); sp.add_argument("--year", type=int)
    sp.set_defaults(fn=cmd_sms)
    sp = sub.add_parser("parse"); sp.add_argument("file"); sp.add_argument("--year", type=int); sp.set_defaults(fn=cmd_parse)
    a = p.parse_args(argv)
    try:
        a.fn(a)
    except BrokenPipeError:
        pass


if __name__ == "__main__":
    main()
